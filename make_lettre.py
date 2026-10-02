from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL

doc = Document()

style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(12)

def add(text, bold=False, align=None, size=12, space_after=8):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p

def ph(placeholder):
    r = doc.add_run(placeholder)
    r.font.highlight_color = None
    return r

# Letterhead
add("TEWVIGH W AL IKHLAS - SARL", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=16, space_after=2)
add("Arafat, Secteur 4A, Lot N° 1575 – Nouakchott, Mauritanie", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
add("Tél : 46 11 46 27   |   NIF : 01730241   |   RCCM : 132231/1595", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
add("Compte bancaire : BPM – 1006080", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
add("", space_after=0)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(24)
pPr = p._p.get_or_add_pPr()
from docx.oxml.ns import qn
pBdr = pPr.makeelement(qn('w:pBdr'), {})
pBdr.set(qn('w:val'), 'single')
pPr.append(pBdr)

# Reference block
add("Nouakchott, le [DATE]", align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=2)
add("Réf : [RÉFÉRENCE]", align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=14)

add("LETTRE DE DÉLÉGATION / POUVOIR", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=14, space_after=14)

body = (
    "Je soussigné(e), [NOM DU GÉRANT / SIGNATAIRE], agissant en qualité de [FONCTION : GÉRANT] "
    "de la société TEWVIGH W AL IKHLAS - SARL, Société à Responsabilité Limitée, immatriculée au "
    "Registre de Commerce sous le numéro 132231/1595 et sous le numéro d'Identification Fiscale (NIF) "
    "01730241, dont le siège social est situé à Arafat, Secteur 4A, Lot N° 1575, Nouakchott – Mauritanie,\n\n"
    "Délègue par la présente, et pour une durée de [DURÉE], Monsieur/Madame [NOM DU DÉLÉGATAIRE], "
    "né(e) le [DATE DE NAISSANCE] à [LIEU], titulaire de la pièce d'identité n° [N° PIÈCE D'IDENTITÉ / NIF], aux fins de :\n\n"
    "1. Représenter la société TEWVIGH W AL IKHLAS - SARL auprès de l'administration des douanes et de "
    "tous services compétents (ASYCUDA, DGD, ANSSA, banques, transitaires, etc.) ;\n\n"
    "2. Déposer, signer et suivre toutes déclarations en douane, demandes de dédouanement, demandes "
    "d'immatriculation et tous documents afférents à l'importation et à l'exportation de marchandises ;\n\n"
    "3. Recevoir et délivrer tous documents, reçus, quittances et correspondances nécessaires ;\n\n"
    "4. Signer tous actes et effectuer toutes démarches relatives à la présente délégation.\n\n"
    "La présente lettre est délivrée pour servir et valoir ce que de droit."
)
for para in body.split("\n\n"):
    add(para, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=8)

add("", space_after=0)
add("Fait à Nouakchott, le [DATE]", space_after=30)
add("Signature et cachet de la société :", space_after=40)
add("______________________________", space_after=6)
add("[NOM DU GÉRANT / SIGNATAIRE]", bold=True, space_after=0)
add("Gérant, TEWVIGH W AL IKHLAS - SARL", space_after=0)

doc.save(r"C:\Users\LMR\Downloads\BLtoXML\Lettre_Delegation.docx")
print("OK")
